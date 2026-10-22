"""
MILESTONE 2 - Document Processing
Load PDF / TXT / DOCX -> clean text -> split into readable chunks.
"""
import os
import re
from pypdf import PdfReader


def load_document(path):
    """Extract text from PDF, TXT, or DOCX."""

    ext = os.path.splitext(path)[1].lower()

    if ext == ".pdf":
        reader = PdfReader(path)
        pages = []

        for page_number, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""

            if page_text.strip():
                pages.append(
                    f"[Page {page_number}]\n{page_text}"
                )

        return "\n\n".join(pages)

    if ext == ".txt":
        with open(path, "r", encoding="utf-8-sig") as file:
            return file.read()

    if ext == ".docx":
        from docx import Document

        doc = Document(path)
        parts = []

        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text)

        # Include tables, which often contain important requirements.
        for table in doc.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                ]

                if any(cells):
                    parts.append(" | ".join(cells))

        return "\n".join(parts)

    raise ValueError(
        "Unsupported file type. Please upload PDF, TXT, or DOCX."
    )


def clean_text(text):
    """Remove extraction artefacts without joining unrelated sentences."""

    if not text:
        return ""

    text = text.replace("\u00ad", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    # Join a word broken across a line by a hyphen.
    text = re.sub(
        r"(?<=[A-Za-z])[-\u2010]\s*\n\s*(?=[a-z])",
        "",
        text,
    )

    # Normalize spaces without removing paragraph boundaries.
    text = re.sub(r"[\t\f\v ]+", " ", text)

    lines = []

    for raw_line in text.split("\n"):
        line = raw_line.strip()

        if not line:
            if lines and lines[-1] != "":
                lines.append("")
            continue

        # Remove standalone page numbers but preserve [Page N].
        if re.fullmatch(r"\d{1,3}", line):
            continue

        lines.append(line)

    text = "\n".join(lines)
    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()


def make_chunks(text, size=850, overlap=100):
    """
    Split text into readable chunks in the original document order.
    Prefer paragraph and sentence boundaries; avoid cutting words.
    """

    text = clean_text(text)

    if not text:
        return []

    if size < 100:
        raise ValueError("Chunk size must be at least 100 characters.")

    overlap = max(0, min(overlap, size // 3))

    # Keep blank-line-separated paragraphs together.
    paragraphs = re.split(r"\n\s*\n+", text)

    chunks = []
    current = ""

    def save_current():
        nonlocal current

        if current.strip():
            chunks.append(current.strip())

        current = ""

    for paragraph in paragraphs:
        paragraph = paragraph.strip()

        if not paragraph:
            continue

        # A normal paragraph fits into a single chunk.
        if len(paragraph) <= size:
            candidate = (
                current + "\n\n" + paragraph
                if current else paragraph
            )

            if len(candidate) <= size:
                current = candidate
            else:
                save_current()
                current = paragraph

            continue

        # A long paragraph needs sentence-based splitting.
        save_current()

        sentences = re.split(
            r"(?<=[.!?])\s+",
            paragraph
        )

        piece = ""

        for sentence in sentences:
            if len(sentence) > size:
                if piece:
                    chunks.append(piece.strip())
                    piece = ""

                start = 0

                while start < len(sentence):
                    end = min(start + size, len(sentence))

                    if end < len(sentence):
                        space = sentence.rfind(" ", start, end)

                        if space > start + size // 2:
                            end = space

                    part = sentence[start:end].strip()

                    if part:
                        chunks.append(part)

                    if end >= len(sentence):
                        break

                    start = max(0, end - overlap)

            elif not piece:
                piece = sentence

            elif len(piece) + len(sentence) + 1 <= size:
                piece += " " + sentence

            else:
                chunks.append(piece.strip())
                piece = sentence

        if piece.strip():
            chunks.append(piece.strip())

    save_current()

    return chunks


def process_file(path):
    """Complete pipeline: extract, clean, and split a document."""

    text = load_document(path)
    cleaned = clean_text(text)

    return make_chunks(cleaned)


if __name__ == "__main__":
    import sys

    if len(sys.argv) > 1:
        file_path = sys.argv[1]
    else:
        file_path = os.path.join(
            os.path.dirname(os.path.abspath(__file__)),
            "data",
            "sample_safety_manual.txt",
        )

    chunks = process_file(file_path)

    print("Total chunks:", len(chunks))

    for index, chunk in enumerate(chunks):
        print(f"\n--- Chunk {index} ---")
        print(chunk)